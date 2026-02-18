"""DOCX parser using python-docx."""
import logging
from typing import Optional

logger = logging.getLogger(__name__)


def parse_docx(file_path: str) -> dict:
    result = {
        "page_count": 1,
        "tables": [],
        "texts": [],
        "confidence": 0.0,
        "errors": [],
    }

    try:
        from docx import Document

        doc = Document(file_path)

        # Extract paragraphs as text
        body_texts = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                body_texts.append(text)

        if body_texts:
            result["texts"].append(
                {
                    "page": 1,
                    "content": "\n".join(body_texts),
                    "type": "body",
                }
            )

        # Extract tables
        for tbl_idx, table in enumerate(doc.tables):
            rows_data = []
            for row in table.rows:
                row_data = [cell.text.strip() for cell in row.cells]
                rows_data.append(row_data)

            if rows_data:
                headers = rows_data[0] if rows_data else []
                data_rows = rows_data[1:] if len(rows_data) > 1 else []
                result["tables"].append(
                    {
                        "page": 1,
                        "index": tbl_idx,
                        "headers": headers,
                        "data": {"columns": headers, "rows": data_rows},
                        "raw_text": str(rows_data),
                        "confidence": 0.85,
                    }
                )

        result["confidence"] = 0.85 if (result["tables"] or result["texts"]) else 0.1

    except ImportError:
        result["errors"].append("python-docx not available")
    except Exception as e:
        result["errors"].append(f"DOCX parse error: {str(e)}")
        logger.error(f"Failed to parse DOCX {file_path}: {e}")

    return result
